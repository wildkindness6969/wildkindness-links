"""Excel export: one workbook with an All Messages sheet and a
Conversations summary sheet (the proto-roster for Converse AI)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..models import ExportBundle, pretty_handle
from .base import Exporter, ProgressFn

EXCEL_MAX_ROWS = 1_048_576

MESSAGE_HEADERS = [
    "Date", "Time", "Direction", "Contact Name", "Contact Number",
    "Message", "Conversation Name", "Group Chat", "Service",
    "Has Attachment", "Conversation ID",
]
MESSAGE_WIDTHS = [12, 10, 10, 24, 18, 80, 24, 10, 10, 13, 38]

SUMMARY_HEADERS = [
    "Conversation Name", "Participants", "Numbers", "Group Chat",
    "Messages", "Sent", "Received", "First Message", "Last Message",
]
SUMMARY_WIDTHS = [28, 32, 32, 10, 10, 8, 10, 18, 18]


class ExcelExporter(Exporter):
    name = "excel"

    def export(self, bundle: ExportBundle, out_dir: Path,
               progress: ProgressFn | None = None) -> list[Path]:
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / "messages.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "All Messages"
        _style_header(ws, MESSAGE_HEADERS, MESSAGE_WIDTHS)

        all_messages = sorted(
            ((m, c) for c in bundle.conversations for m in c.messages),
            key=lambda pair: pair[0].timestamp,
        )
        total = len(all_messages) or 1
        truncated = False
        for i, (msg, convo) in enumerate(all_messages):
            if ws.max_row >= EXCEL_MAX_ROWS:
                truncated = True
                break
            if msg.is_from_me:
                contact_name = convo.display_name
                contact_number = ", ".join(
                    pretty_handle(p.handle) for p in convo.participants)
            else:
                contact_name = msg.sender_name
                contact_number = pretty_handle(msg.sender_handle)
            ws.append([
                msg.timestamp.strftime("%Y-%m-%d"),
                msg.timestamp.strftime("%H:%M:%S"),
                "Sent" if msg.is_from_me else "Received",
                contact_name,
                contact_number,
                msg.text,
                convo.display_name,
                "Yes" if convo.is_group else "No",
                msg.service,
                "Yes" if msg.has_attachments else "No",
                msg.conversation_id,
            ])
            if progress and i % 2000 == 0:
                progress(f"Excel: writing message {i + 1:,} of {total:,}",
                         0.05 + 0.85 * (i / total))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        summary = wb.create_sheet("Conversations")
        _style_header(summary, SUMMARY_HEADERS, SUMMARY_WIDTHS)
        for convo in bundle.conversations:
            summary.append([
                convo.display_name,
                ", ".join(p.display for p in convo.participants),
                ", ".join(pretty_handle(p.handle) for p in convo.participants),
                "Yes" if convo.is_group else "No",
                len(convo.messages),
                convo.sent_count,
                convo.received_count,
                convo.first_timestamp.strftime("%Y-%m-%d %H:%M")
                if convo.first_timestamp else "",
                convo.last_timestamp.strftime("%Y-%m-%d %H:%M")
                if convo.last_timestamp else "",
            ])
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = summary.dimensions

        if truncated:
            bundle.warnings.append(
                "The export exceeds Excel's row limit; the All Messages sheet "
                f"was cut at {EXCEL_MAX_ROWS:,} rows.")

        if progress:
            progress("Excel: saving workbook…", 0.95)
        wb.save(path)
        return [path]


def _style_header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    bold = Font(bold=True)
    for col, width in enumerate(widths, start=1):
        ws.cell(row=1, column=col).font = bold
        ws.column_dimensions[get_column_letter(col)].width = width
