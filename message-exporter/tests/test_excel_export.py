from openpyxl import load_workbook

from exporter.exporters.excel_export import ExcelExporter


def test_excel_export(bundle, tmp_path):
    paths = ExcelExporter().export(bundle, tmp_path)
    assert len(paths) == 1
    wb = load_workbook(paths[0])
    assert wb.sheetnames == ["All Messages", "Conversations"]

    ws = wb["All Messages"]
    headers = [c.value for c in ws[1]]
    assert headers[:6] == ["Date", "Time", "Direction", "Contact Name",
                           "Contact Number", "Message"]
    assert ws.max_row == bundle.total_messages + 1

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    directions = {r[2] for r in rows}
    assert directions == {"Sent", "Received"}
    # interleaved across conversations, sorted by date+time
    keys = [(r[0], r[1]) for r in rows]
    assert keys == sorted(keys)
    # spot-check one received message
    alice_rows = [r for r in rows if r[3] == "Alice Client"]
    assert any("goldendoodle" in r[5] for r in alice_rows)
    assert alice_rows[0][4] == "(925) 555-1234"

    summary = wb["Conversations"]
    assert summary.max_row == len(bundle.conversations) + 1
    by_name = {r[0]: r for r in summary.iter_rows(min_row=2, values_only=True)}
    wedding = by_name["Wedding Crew"]
    assert wedding[3] == "Yes"          # group
    assert wedding[4] == 3              # message count
    assert wedding[5] == 1 and wedding[6] == 2   # sent / received
